
import openpyxl
import json
import sys

def main():
    try:
        file_path = r'c:\Users\MOSTAFA AMIR\Downloads\اسما اصناف خضار - فاكهة (1).xlsx'
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        
        results = {
            "sheet_names": workbook.sheetnames,
            "rows": []
        }
        
        sheet = workbook.active
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i > 50:
                break
            results["rows"].append([str(cell) for cell in row])
        
        print(json.dumps(results, ensure_ascii=False, indent=2))
        
    except Exception as e:
        print(str(e), file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
